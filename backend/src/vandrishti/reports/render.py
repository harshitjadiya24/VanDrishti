"""Render a report dictionary (`reports.data.build`) as Excel (openpyxl) or PDF (reportlab).

Both use only built-in fonts and draw everything locally: no network access.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SUMMARY_LABELS = [
    ("frames", "Frames with a reliable time in the period"),
    ("animal_frames_settled", "Animal frames, settled"),
    ("animal_frames_unsettled", "Animal frames waiting for a person"),
    ("blank_frames", "Likely blank frames"),
    ("people_vehicle_frames", "Frames with people or vehicles (blurred)"),
    ("species", "Species or groups"),
    ("tiger_frames", "Tiger frames"),
    ("named_tigers_seen", "Named tigers seen"),
    ("people_vehicle_events", "People/vehicle events (not dismissed)"),
    ("cameras_operating", "Cameras that recorded frames"),
    ("cameras_registered", "Cameras registered"),
    ("frames_without_reliable_time_all_periods", "Frames with no reliable time (any period)"),
]


def _title(r: dict[str, Any]) -> str:
    p = r["period"]
    span = p["start"] if p["start"] == p["end"] else f"{p['start']} to {p['end']}"
    return f"{r['reserve']['name']} ({r['reserve']['code']}): {p['kind']} report, {span}"


def _tables(r: dict[str, Any]) -> list[tuple[str, list[str], list[list[Any]]]]:
    s = r["summary"]
    return [
        (
            "Summary",
            ["Measure", "Value"],
            [[label, s[k]] for k, label in SUMMARY_LABELS]
            + [
                [f"Risk level {lvl}", s["risk_levels"].get(lvl, 0)]
                for lvl in ("high", "medium", "low")
            ],
        ),
        (
            "Species",
            ["Species", "Frames", "Cameras", "Naive occupancy", "Labelled by a person"],
            [
                [x["label"], x["frames"], x["cameras"], x["naive_occupancy"], x["by_person"]]
                for x in r["species"]
            ],
        ),
        (
            "Tigers",
            ["Tiger", "Sightings", "Cameras", "First", "Last"],
            [
                [t["code"], t["sightings"], ", ".join(t["cameras"]), t["first"], t["last"]]
                for t in r["tigers"]
            ],
        ),
        (
            "Events",
            ["Start (local)", "Camera", "Level", "Score", "People", "Vehicles", "Night", "Reasons"],
            [
                [
                    e["start"],
                    e["camera"],
                    e["level"],
                    e["score"],
                    e["people"],
                    e["vehicles"],
                    "yes" if e["night"] else "no",
                    "; ".join(e["reasons"]),
                ]
                for e in r["events"]
            ],
        ),
        (
            "Cameras",
            [
                "Camera",
                "Status",
                "Frames",
                "Animal frames",
                "Tiger frames",
                "People/vehicle events",
            ],
            [
                [
                    c["code"],
                    c["status"],
                    c["frames"],
                    c["animal_frames"],
                    c["tiger_frames"],
                    c["people_vehicle_events"],
                ]
                for c in r["cameras"]
            ],
        ),
        (
            "Daily",
            ["Date", "Frames", "Animal frames", "Tiger frames", "People/vehicle events"],
            [
                [
                    d["date"],
                    d["frames"],
                    d["animal_frames"],
                    d["tiger_frames"],
                    d["people_vehicle_events"],
                ]
                for d in r["daily"]
            ],
        ),
    ]


def to_xlsx(r: dict[str, Any]) -> bytes:
    wb = Workbook()
    about = wb.active
    assert about is not None  # noqa: S101 (a new workbook always has one sheet)
    about.title = "About"
    about.append([_title(r)])
    about["A1"].font = Font(bold=True, size=13)
    about.append([f"Generated {r['generated_at']} ({r['reserve']['timezone']})"])
    about.append([])
    for c in r["caveats"]:
        about.append([c])
    about.column_dimensions["A"].width = 110
    for name, header, rows in _tables(r):
        ws = wb.create_sheet(name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, start=1):
            longest = max([len(str(h))] + [len(str(row[i - 1] or "")) for row in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 80)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _daily_chart(r: dict[str, Any]) -> Drawing | None:
    days = r["daily"]
    if not any(d["animal_frames"] or d["people_vehicle_events"] for d in days):
        return None
    drawing = Drawing(170 * mm, 60 * mm)
    chart = VerticalBarChart()
    chart.x, chart.y = int(12 * mm), int(10 * mm)
    chart.width, chart.height = int(150 * mm), int(45 * mm)
    chart.data = [
        [d["animal_frames"] for d in days],
        [d["tiger_frames"] for d in days],
        [d["people_vehicle_events"] for d in days],
    ]
    step = max(1, len(days) // 10)
    chart.categoryAxis.categoryNames = [
        d["date"][5:] if i % step == 0 else "" for i, d in enumerate(days)
    ]
    chart.categoryAxis.labels.fontSize = 6
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.bars[0].fillColor = colors.HexColor("#6b94b8")
    chart.bars[1].fillColor = colors.HexColor("#e3a33b")
    chart.bars[2].fillColor = colors.HexColor("#e5585c")
    drawing.add(chart)
    return drawing


def to_pdf(r: dict[str, Any]) -> bytes:
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    heading = styles["Heading2"].clone("heading", keepWithNext=1)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=_title(r),
        author="VanDrishti",
    )
    story: list[Any] = [
        Paragraph(_title(r), styles["Title"]),
        Paragraph(f"Generated {r['generated_at']} ({r['reserve']['timezone']}).", small),
    ]
    for c in r["caveats"]:
        story.append(Paragraph(f"• {c}", small))
    chart = _daily_chart(r)
    if chart is not None:
        story += [
            Spacer(1, 4 * mm),
            Paragraph(
                "Per day: animal frames (blue), tiger frames (amber), people/vehicle events (red)",
                small,
            ),
            chart,
        ]
    style = TableStyle(
        [
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ece9")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2bd")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
    )
    for name, header, rows in _tables(r):
        if name == "Daily":
            continue  # the chart shows it; the Excel file has the table
        story += [Spacer(1, 5 * mm), Paragraph(name, heading)]
        if not rows:
            story.append(Paragraph("None in this period.", small))
            continue
        cells = [header] + [
            [
                Paragraph(str(v), small)
                if isinstance(v, str) and len(v) > 40
                else ("" if v is None else v)
                for v in row
            ]
            for row in rows[:300]
        ]
        story.append(Table(cells, repeatRows=1, style=style, hAlign="LEFT"))
        if len(rows) > 300:
            story.append(
                Paragraph(f"First 300 of {len(rows)} rows; the Excel report has all.", small)
            )
    doc.build(story)
    return buf.getvalue()
